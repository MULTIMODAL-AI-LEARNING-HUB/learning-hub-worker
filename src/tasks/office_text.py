"""Office/plain-text extraction for document processing tasks."""

import io
import logging

logger = logging.getLogger("worker.office_text")

_MAX_CHARS = 500_000


def extract_text_from_office_file(file_bytes: bytes, ext: str) -> str:
    """Extract readable text from txt/doc/docx bytes.

    python-docx is already a worker dependency; legacy .doc (OLE) is parsed
    best-effort by stripping binary noise.
    """
    ext = (ext or "").lower().strip(".")
    try:
        if ext == "txt":
            return _truncate(_decode_text(file_bytes))
        if ext == "docx":
            return _truncate(_extract_docx(file_bytes))
        if ext == "doc":
            return _truncate(_extract_legacy_doc(file_bytes))
    except Exception as exc:  # never crash a task on extraction
        logger.warning("Office text extraction failed for .%s: %s", ext, exc)
    return ""


def _decode_text(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except (UnicodeDecodeError, ValueError):
            continue
    return raw.decode("utf-8", errors="ignore")


def _extract_docx(raw: bytes) -> str:
    try:
        from docx import Document
    except ImportError:
        logger.warning("python-docx not installed; cannot parse .docx")
        return ""
    doc = Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text and cell.text.strip():
                    parts.append(cell.text.strip())
    return "\n".join(parts)


def _extract_legacy_doc(raw: bytes) -> str:
    # Legacy .doc is an OLE compound binary file. Best-effort: decode as text
    # and keep long printable runs. Users should prefer .docx, but this avoids
    # marking the document as failed when text is recoverable.
    text = _decode_text(raw)
    cleaned = "".join(ch if ch.isprintable() or ch in "\n\t" else " " for ch in text)
    runs = [run.strip() for run in cleaned.split("  ") if run.strip()]
    long_runs = [run for run in runs if len(run) >= 8]
    return "\n".join(long_runs)


def _truncate(text: str) -> str:
    text = (text or "").strip()
    if len(text) > _MAX_CHARS:
        return text[:_MAX_CHARS]
    return text
